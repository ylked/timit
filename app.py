import csv
import io
import json
import os
import sqlite3
from calendar import monthrange
from datetime import date, datetime, timedelta
from pathlib import Path

from flask import Flask, g, jsonify, render_template, request, send_file

BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "data" / "timit.db"

app = Flask(__name__)


def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_PATH)
        g.db.row_factory = sqlite3.Row
    return g.db


@app.teardown_appcontext
def close_db(exception=None):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def init_db():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(DB_PATH)
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS entries (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            start_ts TEXT NOT NULL,
            end_ts TEXT,
            note TEXT DEFAULT ''
        )
        """
    )
    db.commit()
    db.close()


# ---------- helpers ----------

def now_iso():
    return datetime.now().replace(microsecond=0).isoformat()


def parse_iso(s):
    return datetime.fromisoformat(s)


def normalize_iso(s):
    return datetime.fromisoformat(s).replace(microsecond=0).isoformat()


def effective_end(row):
    if row["end_ts"] is None:
        return datetime.now().replace(microsecond=0)
    return parse_iso(row["end_ts"])


def seconds_to_hhmm(secs):
    secs = int(round(secs))
    h, rem = divmod(secs, 3600)
    m, _ = divmod(rem, 60)
    return f"{h:02d}:{m:02d}"


def overlapping_entries(db, range_start, range_end):
    return db.execute(
        "SELECT * FROM entries WHERE start_ts < ? AND (end_ts IS NULL OR end_ts > ?) ORDER BY start_ts",
        (range_end.isoformat(), range_start.isoformat()),
    ).fetchall()


def split_by_day(start_dt, end_dt):
    cur = start_dt
    while cur < end_dt:
        day_end = datetime.combine(cur.date() + timedelta(days=1), datetime.min.time())
        chunk_end = min(end_dt, day_end)
        yield cur.date(), (chunk_end - cur).total_seconds()
        cur = chunk_end


def daily_totals(db, start_date, end_date):
    range_start = datetime.combine(start_date, datetime.min.time())
    range_end = datetime.combine(end_date + timedelta(days=1), datetime.min.time())
    rows = overlapping_entries(db, range_start, range_end)
    totals = {}
    for row in rows:
        s = max(parse_iso(row["start_ts"]), range_start)
        e = min(effective_end(row), range_end)
        for d, secs in split_by_day(s, e):
            totals[d] = totals.get(d, 0) + secs
    return totals


def resolve_range(db, start_s, end_s):
    if start_s and end_s:
        return date.fromisoformat(start_s), date.fromisoformat(end_s)
    row = db.execute("SELECT MIN(start_ts) AS mn, MAX(start_ts) AS mx FROM entries").fetchone()
    if row["mn"] is None:
        today = date.today()
        return today, today
    mn = parse_iso(row["mn"]).date()
    mx = parse_iso(row["mx"]).date()
    return (date.fromisoformat(start_s) if start_s else mn, date.fromisoformat(end_s) if end_s else mx)


def get_entries_in_range(db, start_date, end_date):
    range_start = datetime.combine(start_date, datetime.min.time())
    range_end = datetime.combine(end_date + timedelta(days=1), datetime.min.time())
    return overlapping_entries(db, range_start, range_end)


# ---------- pages ----------

@app.get("/")
def index():
    return render_template("index.html")


# ---------- timer ----------

@app.get("/api/timer/status")
def timer_status():
    db = get_db()
    row = db.execute("SELECT * FROM entries WHERE end_ts IS NULL ORDER BY id DESC LIMIT 1").fetchone()
    if row is None:
        return jsonify({"running": False})
    return jsonify({"running": True, "entry": dict(row)})


@app.post("/api/timer/start")
def timer_start():
    db = get_db()
    if db.execute("SELECT id FROM entries WHERE end_ts IS NULL").fetchone() is not None:
        return jsonify({"error": "A timer is already running"}), 409
    note = (request.get_json(silent=True) or {}).get("note", "")
    cur = db.execute("INSERT INTO entries (start_ts, end_ts, note) VALUES (?, NULL, ?)", (now_iso(), note))
    db.commit()
    row = db.execute("SELECT * FROM entries WHERE id = ?", (cur.lastrowid,)).fetchone()
    return jsonify(dict(row))


@app.post("/api/timer/stop")
def timer_stop():
    db = get_db()
    row = db.execute("SELECT * FROM entries WHERE end_ts IS NULL ORDER BY id DESC LIMIT 1").fetchone()
    if row is None:
        return jsonify({"error": "No timer is running"}), 409
    db.execute("UPDATE entries SET end_ts = ? WHERE id = ?", (now_iso(), row["id"]))
    db.commit()
    row = db.execute("SELECT * FROM entries WHERE id = ?", (row["id"],)).fetchone()
    return jsonify(dict(row))


# ---------- entries (CRUD) ----------

@app.get("/api/entries")
def list_entries():
    start_s = request.args.get("start")
    end_s = request.args.get("end")
    db = get_db()
    if start_s or end_s:
        start_date, end_date = resolve_range(db, start_s, end_s)
        rows = get_entries_in_range(db, start_date, end_date)
    else:
        rows = db.execute("SELECT * FROM entries ORDER BY start_ts DESC").fetchall()
    return jsonify([dict(r) for r in rows])


@app.post("/api/entries")
def create_entry():
    data = request.get_json()
    start_ts = normalize_iso(data["start_ts"])
    end_ts = normalize_iso(data["end_ts"]) if data.get("end_ts") else None
    note = data.get("note", "")
    if end_ts and end_ts <= start_ts:
        return jsonify({"error": "End must be after start"}), 400
    db = get_db()
    if end_ts is None and db.execute("SELECT id FROM entries WHERE end_ts IS NULL").fetchone() is not None:
        return jsonify({"error": "A timer is already running"}), 409
    cur = db.execute("INSERT INTO entries (start_ts, end_ts, note) VALUES (?, ?, ?)", (start_ts, end_ts, note))
    db.commit()
    row = db.execute("SELECT * FROM entries WHERE id = ?", (cur.lastrowid,)).fetchone()
    return jsonify(dict(row))


@app.put("/api/entries/<int:entry_id>")
def update_entry(entry_id):
    data = request.get_json()
    db = get_db()
    row = db.execute("SELECT * FROM entries WHERE id = ?", (entry_id,)).fetchone()
    if row is None:
        return jsonify({"error": "Not found"}), 404
    start_ts = normalize_iso(data["start_ts"]) if data.get("start_ts") else row["start_ts"]
    end_ts = (normalize_iso(data["end_ts"]) if data.get("end_ts") else None) if "end_ts" in data else row["end_ts"]
    note = data.get("note", row["note"])
    if end_ts and end_ts <= start_ts:
        return jsonify({"error": "End must be after start"}), 400
    db.execute("UPDATE entries SET start_ts=?, end_ts=?, note=? WHERE id=?", (start_ts, end_ts, note, entry_id))
    db.commit()
    row = db.execute("SELECT * FROM entries WHERE id = ?", (entry_id,)).fetchone()
    return jsonify(dict(row))


@app.delete("/api/entries/<int:entry_id>")
def delete_entry(entry_id):
    db = get_db()
    db.execute("DELETE FROM entries WHERE id = ?", (entry_id,))
    db.commit()
    return jsonify({"ok": True})


# ---------- summaries ----------

@app.get("/api/summary/day")
def summary_day():
    db = get_db()
    start_date, end_date = resolve_range(db, request.args.get("start"), request.args.get("end"))
    totals = daily_totals(db, start_date, end_date)
    result = []
    d = start_date
    while d <= end_date:
        result.append({"date": d.isoformat(), "seconds": totals.get(d, 0)})
        d += timedelta(days=1)
    return jsonify(result)


@app.get("/api/summary/week")
def summary_week():
    db = get_db()
    start_date, end_date = resolve_range(db, request.args.get("start"), request.args.get("end"))
    totals = daily_totals(db, start_date, end_date)
    weeks = {}
    d = start_date
    while d <= end_date:
        week_start = d - timedelta(days=d.weekday())
        weeks[week_start] = weeks.get(week_start, 0) + totals.get(d, 0)
        d += timedelta(days=1)
    result = [{"week_start": ws.isoformat(), "seconds": secs} for ws, secs in sorted(weeks.items())]
    return jsonify(result)


@app.get("/api/calendar")
def calendar_view():
    year = int(request.args.get("year"))
    month = int(request.args.get("month"))
    start_date = date(year, month, 1)
    end_date = date(year, month, monthrange(year, month)[1])
    db = get_db()
    totals = daily_totals(db, start_date, end_date)
    return jsonify({d.isoformat(): secs for d, secs in totals.items()})


# ---------- exports ----------

@app.get("/api/export/xlsx")
def export_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font

    db = get_db()
    start_date, end_date = resolve_range(db, request.args.get("start"), request.args.get("end"))
    totals = daily_totals(db, start_date, end_date)
    rows = get_entries_in_range(db, start_date, end_date)

    entry_counts = {}
    for row in rows:
        d = parse_iso(row["start_ts"]).date()
        entry_counts[d] = entry_counts.get(d, 0) + 1

    wb = Workbook()
    ws = wb.active
    ws.title = "Daily Summary"
    header_font = Font(bold=True)
    ws.append(["Date", "Weekday", "Total Hours", "Total (HH:MM)", "Entries"])
    for cell in ws[1]:
        cell.font = header_font

    d = start_date
    total_seconds = 0
    total_entries = 0
    while d <= end_date:
        secs = totals.get(d, 0)
        total_seconds += secs
        total_entries += entry_counts.get(d, 0)
        ws.append([d.isoformat(), d.strftime("%A"), round(secs / 3600, 2), seconds_to_hhmm(secs), entry_counts.get(d, 0)])
        d += timedelta(days=1)

    ws.append([])
    ws.append(["Total", "", round(total_seconds / 3600, 2), seconds_to_hhmm(total_seconds), total_entries])
    for cell in ws[ws.max_row]:
        cell.font = header_font
    for col, width in zip("ABCDE", [12, 12, 12, 14, 10]):
        ws.column_dimensions[col].width = width

    ws2 = wb.create_sheet("Raw Entries")
    ws2.append(["ID", "Start", "End", "Duration (HH:MM)", "Note"])
    for cell in ws2[1]:
        cell.font = header_font
    for row in rows:
        s = parse_iso(row["start_ts"])
        e = effective_end(row)
        dur = (e - s).total_seconds()
        ws2.append([row["id"], row["start_ts"], row["end_ts"] or "(running)", seconds_to_hhmm(dur), row["note"] or ""])
    for col, width in zip("ABCDE", [6, 20, 20, 16, 30]):
        ws2.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"timit-report-{start_date}-to-{end_date}.xlsx"
    return send_file(
        buf,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.get("/api/export/csv")
def export_csv():
    db = get_db()
    start_s, end_s = request.args.get("start"), request.args.get("end")
    if start_s or end_s:
        start_date, end_date = resolve_range(db, start_s, end_s)
        rows = get_entries_in_range(db, start_date, end_date)
    else:
        rows = db.execute("SELECT * FROM entries ORDER BY start_ts").fetchall()

    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["id", "start_ts", "end_ts", "duration_seconds", "note"])
    for row in rows:
        dur = int((effective_end(row) - parse_iso(row["start_ts"])).total_seconds()) if row["end_ts"] else ""
        writer.writerow([row["id"], row["start_ts"], row["end_ts"] or "", dur, row["note"] or ""])
    mem = io.BytesIO(buf.getvalue().encode("utf-8"))
    return send_file(mem, as_attachment=True, download_name="timit-entries.csv", mimetype="text/csv")


@app.get("/api/export/json")
def export_json():
    db = get_db()
    start_s, end_s = request.args.get("start"), request.args.get("end")
    if start_s or end_s:
        start_date, end_date = resolve_range(db, start_s, end_s)
        rows = get_entries_in_range(db, start_date, end_date)
    else:
        rows = db.execute("SELECT * FROM entries ORDER BY start_ts").fetchall()
    data = [dict(r) for r in rows]
    mem = io.BytesIO(json.dumps(data, indent=2).encode("utf-8"))
    return send_file(mem, as_attachment=True, download_name="timit-entries.json", mimetype="application/json")


@app.get("/api/backup")
def backup():
    return send_file(
        DB_PATH,
        as_attachment=True,
        download_name=f"timit-backup-{datetime.now():%Y%m%d-%H%M%S}.db",
        mimetype="application/x-sqlite3",
    )


init_db()

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5252))
    app.run(debug=True, port=port)
